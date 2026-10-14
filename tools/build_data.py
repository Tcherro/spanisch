#!/usr/bin/env python3
"""Baut data/vocab.json aus der Master-Excel (Sheet "Vokabeln" + "Redemittel").

Aufruf: python3 tools/build_data.py <pfad/zur/Vokabelliste.xlsx> [--date YYYY-MM-DD]

- IDs sind stabil (Hash aus Spanisch|Deutsch), damit der Lernfortschritt
  auf dem iPhone bei Updates erhalten bleibt.
- Das Feld "added" bleibt für bekannte Einträge erhalten; neue Einträge
  bekommen das Lauf-Datum und erscheinen in der App unter "Neu diese Woche".
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "vocab.json"
META = ROOT / "data" / "meta.json"


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def vid(es, de):
    return hashlib.sha1(f"{es.lower()}|{de.lower()}".encode()).hexdigest()[:10]


def lesson_of(src):
    m = re.match(r"\s*L(\d+)", src or "")
    return int(m.group(1)) if m else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--date", default=dt.date.today().isoformat())
    a = ap.parse_args()

    old = {}
    if OUT.exists():
        for v in json.loads(OUT.read_text(encoding="utf-8"))["items"]:
            old[v["id"]] = v

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    items, seen = [], set()

    def add(es, de, src, note=""):
        es, de, src = clean(es), clean(de), clean(src)
        if not es or not de:
            return
        i = vid(es, de)
        if i in seen:
            return
        seen.add(i)
        prev = old.get(i)
        if prev:
            added = prev.get("added")
        elif old:  # echter Neuzugang seit dem letzten Lauf
            added = a.date
        else:  # Erstaufbau: nur die letzten Lektionen als "neu" markieren
            l = lesson_of(src)
            added = {28: "2026-10-14", 27: "2026-10-12"}.get(l)
        it = {"id": i, "es": es, "de": de, "l": lesson_of(src), "src": src}
        if note:
            it["n"] = clean(note)
        if added:
            it["added"] = added
        items.append(it)

    for r in list(wb["Vokabeln"].iter_rows(values_only=True))[1:]:
        add(r[0], r[1], r[2] if len(r) > 2 else "")

    if "Redemittel" in wb.sheetnames:
        for r in wb["Redemittel"].iter_rows(values_only=True):
            if r and r[0] and r[1] and r[2] and r[0] not in ("Kategorie",) and len(r) > 3:
                add(r[1], r[2], "Redemittel – " + clean(r[0]), r[3] if r[3] not in (None, "—") else "")

    OUT.write_text(json.dumps({"generated": a.date, "items": items}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    meta = json.loads(META.read_text(encoding="utf-8")) if META.exists() else {}
    new_count = sum(1 for v in items if v.get("added") == a.date)
    meta.update({
        "version": a.date + "-" + hashlib.sha1(OUT.read_bytes()).hexdigest()[:6],
        "generated": a.date,
        "vocabCount": len(items),
        "maxLesson": max(v["l"] for v in items),
    })
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(items)} Vokabeln, davon {new_count} neu am {a.date}; Version {meta['version']}")


if __name__ == "__main__":
    sys.exit(main())
