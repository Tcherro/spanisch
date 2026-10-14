#!/usr/bin/env python3
"""Übungen aus den Excel-Sheets "Übungen" und "Übungen Position".

  python3 tools/excel_exercises.py <Vokabelliste.xlsx>   → schreibt tools/excel_exercises.json
grammar_src.py liest diese JSON (das Excel selbst liegt nicht im öffentlichen Repo).
"""
import json, re, sys
from pathlib import Path

J = Path(__file__).resolve().parent / "excel_exercises.json"
GAP = re.compile(r"_{2,}|\.\.\.|…")

POOLS = {
    "dem": ["este", "esta", "estos", "estas", "ese", "esa", "esos", "esas", "aquel", "aquella"],
    "pos_ort": ["al lado", "debajo", "a la derecha", "a la izquierda", "entre", "detrás", "encima", "delante", "dentro", "fuera"],
    "contr": ["al", "del", "de la", "de los", "de las"],
    "prep": ["Adónde", "a", "De dónde", "de", "Dónde", "en"],
}


def gap_answer(task, sol):
    parts = GAP.split(task)
    if len(parts) != 2:
        return None
    pre, post = parts[0], parts[1]
    s = sol
    if not s.lower().startswith(pre.lower().strip()[:len(pre.strip())]):
        return None
    a = s[len(pre):]
    if post.strip() and a.endswith(post):
        a = a[: len(a) - len(post)]
    elif post.strip():
        # Satzzeichen-Toleranz
        ps = post.rstrip(".?!")
        i = a.rfind(ps) if ps else -1
        if i < 0:
            return None
        a = a[:i]
    a = a.strip().strip(",")
    return a or None


def build(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out = {"dem": [], "adj": [], "prep": [], "pos_ort": []}

    def mcitem(task, sol, de, ex, pool):
        a = gap_answer(task, sol)
        if not a:
            return None
        cap = a[:1].isupper()
        import random
        rnd = random.Random(task)
        rest = [o for o in pool if o.lower() != a.lower()]
        opts = [a] + rnd.sample(rest, min(3, len(rest)))
        if cap:
            opts = [o[:1].upper() + o[1:] for o in opts]
        return {"t": "mc", "q": GAP.sub("___", task), "de": de or "", "opts": opts, "ex": ex or ""}

    rows = list(wb["Übungen"].iter_rows(values_only=True))
    for r in rows:
        if not r or not isinstance(r[0], int):
            continue
        task, sol = str(r[1]), str(r[2])
        if "..." in task and task.endswith("..."):  # Adjektivendungen → Eingabe
            stem = task[:-3].split()[-1]
            word = [w for w in re.findall(r"[\wáéíóúñ]+", sol) if w.lower().startswith(stem.lower())]
            if word:
                out["adj"].append({"t": "type", "q": task.replace("...", "___"), "de": r[3] or "", "ans": [word[-1], word[-1][len(stem):]],
                                   "ex": word[-1]})
            continue
        pool = POOLS["prep"] if ("dónde" in sol.lower() or " va a " in sol or "Viene" in sol or "Están en" in sol or sol.startswith("Va a")) else POOLS["dem"]
        it = mcitem(task, sol, r[4] if pool is POOLS["dem"] else r[3], r[3] if pool is POOLS["dem"] else r[4], pool)
        if it:
            out["prep" if pool is POOLS["prep"] else "dem"].append(it)

    sec = None
    for r in wb["Übungen Position"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0][:2] in ("A.", "B.", "C.", "D.", "E."):
            sec = r[0][0]
            continue
        if not r or not isinstance(r[0], int):
            continue
        if sec == "A":
            it = mcitem(str(r[1]), str(r[2]), r[4], r[3], POOLS["pos_ort"])
            if it: out["pos_ort"].append(it)
        elif sec == "C":
            it = mcitem(str(r[1]), str(r[2]), r[4], r[3], POOLS["contr"])
            if it: out["pos_ort"].append(it)
        elif sec == "B":
            out["pos_ort"].append({"t": "order", "de": f"Gegenteil von: {r[1]}", "ans": str(r[2]), "ex": r[3] or ""})
        elif sec == "D":
            out["pos_ort"].append({"t": "order", "de": str(r[1]), "ans": str(r[2]), "ex": r[3] or ""})
    J.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print({k: len(v) for k, v in out.items()})


def add_from_excel(topics):
    if not J.exists():
        return
    data = json.loads(J.read_text(encoding="utf-8"))
    by = {t["id"]: t for t in topics}
    for k, items in data.items():
        if k in by:
            by[k]["items"].extend(items)


if __name__ == "__main__":
    build(sys.argv[1])
